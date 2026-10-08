import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)          # #0F172A
    PRIMARY_BLUE = RGBColor(30, 58, 138)     # #1E3A8A
    ACCENT_TEAL = RGBColor(2, 132, 199)      # #0284C7
    AMBER_FIREBASE = RGBColor(245, 158, 11)  # #F59E0B
    CARD_BG = RGBColor(248, 250, 252)        # #F8FAFC
    CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0
    TEXT_DARK = RGBColor(30, 41, 59)         # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B
    TEXT_WHITE = RGBColor(255, 255, 255)
    GREEN_ACCENT = RGBColor(16, 185, 129)    # #10B981

    def add_header(slide, title, category="PHÂN TÍCH & TÍCH HỢP CLOUD / FIREBASE"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        r_cat = p_cat.add_run()
        r_cat.text = category.upper()
        r_cat.font.name = "Arial"
        r_cat.font.size = Pt(10)
        r_cat.font.bold = True
        r_cat.font.color.rgb = ACCENT_TEAL

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        r_title = p_title.add_run()
        r_title.text = title
        r_title.font.name = "Arial"
        r_title.font.size = Pt(22)
        r_title.font.bold = True
        r_title.font.color.rgb = PRIMARY_BLUE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    card_t = add_card(slide1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), bg_color=RGBColor(30, 41, 59), border_color=ACCENT_TEAL)

    tb = slide1.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "HỌC PHẦN: PHÁT TRIỂN ỨNG DỤNG DI ĐỘNG (CSE441) — BÀI TẬP NHÓM"
    r0.font.name = "Arial"
    r0.font.size = Pt(12)
    r0.font.bold = True
    r0.font.color.rgb = AMBER_FIREBASE

    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = "PHÂN TÍCH VÀ LẬP PHƯƠNG ÁN TÍCH HỢP CLOUD\nCHO HỆ THỐNG QUẢN LÝ TÀI LIỆU (DMS)"
    r1.font.name = "Arial"
    r1.font.size = Pt(28)
    r1.font.bold = True
    r1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(14)
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Nghiên cứu Hệ sinh thái Cloud Storage & Hướng dẫn Tích hợp Firebase (Google Auth & Storage) cho Flutter"
    r2.font.name = "Arial"
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = "ĐƠN VỊ THỰC HIỆN: NHÓM 16 — LỚP 65KTPM — KHOA CÔNG NGHỆ THÔNG TIN — ĐẠI HỌC THỦY LỢI (TLU)"
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    r3.font.bold = True
    r3.font.color.rgb = ACCENT_TEAL

    p4 = tf.add_paragraph()
    r4 = p4.add_run()
    r4.text = "Thành viên: Nguyễn Văn Huỳnh (NT) | Lê Anh Tuấn | Nguyễn Trung Kiên | Trần Anh Tuấn"
    r4.font.name = "Arial"
    r4.font.size = Pt(11)
    r4.font.color.rgb = TEXT_WHITE
    p4.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: 7 Checklist Requirements
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "MỤC TIÊU & ĐỐI SOÁT CHECKLIST 7 MỤC ĐỀ BÀI")

    checklists = [
        ("Mục 1", "Phân tích 4 thành phần cốt lõi: Frontend, Backend API, Metadata DB, File Storage."),
        ("Mục 2", "Chỉ rõ 5 điểm nghẽn hạ tầng truyền thống: I/O Disk, Scale-up, SPOF, Ransomware, Chi phí."),
        ("Mục 3", "Lựa chọn mô hình Cloud (Public Cloud) & dịch vụ đối sánh (AWS S3, Azure, Google Cloud)."),
        ("Mục 4", "Thiết kế sơ đồ kiến trúc Cloud và luồng dữ liệu Direct Upload / Streaming."),
        ("Mục 5", "Đánh giá 3 trụ cột tác động: An toàn Bảo mật, Chi phí vận hành (TCO), Hiệu suất."),
        ("Mục 6", "Khai phá giải pháp Firebase: Google Sign-In & Firebase Storage cho ứng dụng Flutter."),
        ("Mục 7", "Tạo bộ Slide báo cáo, cập nhật README.md đồ án và kế hoạch phân công bài tập nhóm.")
    ]

    for idx, (m, desc) in enumerate(checklists):
        row = idx // 2
        col = idx % 2
        if idx == 6:
            c_left = Inches(0.8)
            c_top = Inches(1.6 + 3 * 1.35)
            c_w = Inches(11.733)
            c_h = Inches(1.2)
        else:
            c_left = Inches(0.8 + col * 5.95)
            c_top = Inches(1.6 + row * 1.35)
            c_w = Inches(5.75)
            c_h = Inches(1.2)
        add_card(slide2, c_left, c_top, c_w, c_h)

        tb = slide2.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.1), c_w - Inches(0.3), c_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r_m = p.add_run()
        r_m.text = f"✅ {m}: "
        r_m.font.name = "Arial"
        r_m.font.size = Pt(11)
        r_m.font.bold = True
        r_m.font.color.rgb = PRIMARY_BLUE

        run_desc = p.add_run()
        run_desc.text = desc
        run_desc.font.name = "Arial"
        run_desc.font.size = Pt(10)
        run_desc.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 3: Core Components of DMS
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "1. PHÂN TÍCH 4 THÀNH PHẦN CỐT LÕI CỦA HỆ THỐNG DMS")

    comps = [
        ("TẦNG 1: FRONTEND", "Web SPA & Mobile Flutter", "• Giao diện tra cứu Live Search không dấu.\n• Danh mục môn học, tags, phân loại.\n• Trình đọc tài liệu PDF inline, audio/video.\n• Kéo-thả tải lên kèm thanh tiến trình.\n👉 Tính sẵn sàng Cloud: 95% (Static Asset)."),
        ("TẦNG 2: BACKEND API", "Node.js / FastAPI / Spring Boot", "• Xác thực JWT & phân quyền vai trò (RBAC).\n• Tiếp nhận tệp, quét virus, kiểm tra MIME.\n• Tác vụ nền: Trích xuất OCR, sinh Thumbnail.\n• Quản lý phiên bản & Audit Trail.\n👉 Tính sẵn sàng Cloud: 85% (Stateless API)."),
        ("TẦNG 3: METADATA DB", "PostgreSQL / MySQL / SQLite", "• Lưu trữ bảng documents, subjects.\n• Quản lý kích thước, MIME, Checksum SHA-256.\n• Bảng delete_logs phục vụ kiểm toán đồng bộ.\n• Quan hệ toàn vẹn dữ liệu (ACID).\n👉 Tính sẵn sàng Cloud: 90% (Managed RDS)."),
        ("TẦNG 4: FILE STORAGE", "Local Disk / Mạng NAS / SAN", "• Lưu tệp nhị phân trên đĩa vật lý máy chủ.\n• Backend phải mở luồng đọc/ghi trực tiếp.\n• Cấu trúc thư mục: /uploads/{year}/{doc}.pdf.\n• Nguy cơ quá tải dung lượng và I/O đĩa.\n👉 Tính sẵn sàng: Cần chuyển sang S3/Cloud Storage.")
    ]

    for idx, (title, sub, details) in enumerate(comps):
        c_left = Inches(0.8 + idx * 2.95)
        c_top = Inches(1.6)
        c_w = Inches(2.8)
        c_h = Inches(5.3)
        add_card(slide3, c_left, c_top, c_w, c_h)

        tb = slide3.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.15), c_w - Inches(0.3), c_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r_t = p1.add_run()
        r_t.text = title
        r_t.font.name = "Arial"
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        r_sub = p2.add_run()
        r_sub.text = sub
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(9.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = ACCENT_TEAL
        p2.space_after = Pt(8)

        p3 = tf.add_paragraph()
        r_det = p3.add_run()
        r_det.text = details
        r_det.font.name = "Arial"
        r_det.font.size = Pt(9)
        r_det.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 4: Bottlenecks of Traditional Infrastructure
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "2. NĂM ĐIỂM NGHẼN CỦA HẠ TẦNG TRUYỀN THỐNG (ON-PREMISES)")

    bottlenecks = [
        ("1. GIỚI HẠN DUNG LƯỢNG & NGHẼN I/O", "Đĩa cứng vật lý có dung lượng hữu hạn. Khi hàng nghìn sinh viên tải tài liệu cùng lúc mùa thi, Disk I/O chạm ngưỡng 100%, gây thắt cổ chai toàn hệ thống."),
        ("2. KHÓ KHĂN KHI MỞ RỘNG (SCALE-UP)", "Chỉ có thể nâng cấp đĩa to hơn hoặc RAM lớn hơn (Scale-up) với chi phí cấp số nhân. Không thể tự động mở rộng theo thời gian thực (Auto-scaling)."),
        ("3. ĐIỂM CHẾT DUY NHẤT (SPOF) & VPN", "Toàn bộ dịch vụ nằm tại một phòng máy chủ. Hỏng hóc vật lý sẽ sập toàn bộ. Truy cập từ xa bắt buộc qua VPN phức tạp, tốc độ chậm và hay gián đoạn."),
        ("4. NGUY CƠ THẢM HỌA & RANSOMWARE", "Tệp lưu đường dẫn tĩnh trên máy chủ dễ bị mã hóa Ransomware. Khâu sao lưu băng từ/đĩa ngoài thủ công, RPO và RTO kéo dài từ nhiều giờ đến nhiều ngày."),
        ("5. GÁNH NẶNG CHI PHÍ CAPEX & OPEX", "Chi phí mua sắm phần cứng máy chủ ban đầu cực lớn. Chi phí ẩn khổng lồ: điện điều hòa 24/7, phòng server, khấu hao, nhân sự IT bảo trì thiết bị.")
    ]

    for idx, (title, desc) in enumerate(bottlenecks):
        c_top = Inches(1.6 + idx * 1.05)
        add_card(slide4, Inches(0.8), c_top, Inches(11.733), Inches(0.95))

        tb = slide4.shapes.add_textbox(Inches(1.0), c_top + Inches(0.08), Inches(11.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = f"⚠️ {title}: "
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(220, 38, 38)

        run_d = p1.add_run()
        run_d.text = desc
        run_d.font.name = "Arial"
        run_d.font.size = Pt(9.5)
        run_d.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 5: Cloud Deployment Model & Service Selection
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "3. LỰA CHỌN MÔ HÌNH CLOUD & DỊCH VỤ ĐỐI SÁNH")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.75), Inches(5.3))
    tb_l = slide5.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.35), Inches(5.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    r_l = p.add_run()
    r_l.text = "LỰA CHỌN: PUBLIC CLOUD (ĐÁM MÂY CÔNG CỘNG)"
    r_l.font.name = "Arial"
    r_l.font.size = Pt(12)
    r_l.font.bold = True
    r_l.font.color.rgb = PRIMARY_BLUE

    p_sub = tf_l.add_paragraph()
    r_sub = p_sub.add_run()
    r_sub.text = "• Lý do lựa chọn: Tài liệu học tập không thuộc dữ liệu mật quốc gia, cần phân phối diện rộng cho hàng nghìn sinh viên với chi phí tối ưu nhất.\n• Loại bỏ 100% chi phí đầu tư phòng máy và máy chủ vật lý ban đầu.\n• Sẵn sàng CDN toàn cầu đưa tài liệu về gần sinh viên nhất.\n• Độ bền bỉ dữ liệu đạt chuẩn 99.999999999% (11 số 9)."
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = TEXT_DARK
    p_sub.space_before = Pt(8)

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.3))
    tb_r = slide5.shapes.add_textbox(Inches(7.0), Inches(1.75), Inches(5.35), Inches(5.0))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_r = tf_r.paragraphs[0]
    r_rh = p_r.add_run()
    r_rh.text = "SO SÁNH CÁC DỊCH VỤ CLOUD OBJECT STORAGE"
    r_rh.font.name = "Arial"
    r_rh.font.size = Pt(12)
    r_rh.font.bold = True
    r_rh.font.color.rgb = PRIMARY_BLUE

    comps_cloud = [
        ("Amazon S3 (AWS):", "Chuẩn mực công nghiệp, đầy đủ phân tầng vòng đời (Standard, IA, Glacier Deep Archive tiết kiệm 90%), Pre-signed URL cực mạnh."),
        ("Google Cloud Storage / Firebase:", "Tối ưu hóa đặc biệt cho ứng dụng Flutter và Mobile qua Firebase SDK, tích hợp tự nhiên với Google Sign-In, thiết lập cực nhanh."),
        ("Azure Blob Storage:", "Thích hợp cho doanh nghiệp tích hợp sâu hệ sinh thái Microsoft Office 365 và Active Directory.")
    ]
    for c_title, c_desc in comps_cloud:
        p_c = tf_r.add_paragraph()
        r_ct = p_c.add_run()
        r_ct.text = f"🔹 {c_title} "
        r_ct.font.name = "Arial"
        r_ct.font.size = Pt(10)
        r_ct.font.bold = True
        r_ct.font.color.rgb = ACCENT_TEAL
        p_c.space_before = Pt(8)

        run = p_c.add_run()
        run.text = c_desc
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 6: Cloud Architecture & Data Flow
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "4. SƠ ĐỒ KIẾN TRÚC CLOUD & LUỒNG DIRECT UPLOAD")

    arch_img_path = os.path.join("scripts", "output", "cloud_dms_architecture.png")
    if os.path.exists(arch_img_path):
        slide6.shapes.add_picture(arch_img_path, Inches(0.8), Inches(1.6), width=Inches(7.2))
    else:
        add_card(slide6, Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.3))

    add_card(slide6, Inches(8.2), Inches(1.6), Inches(4.333), Inches(5.3))
    tb_flow = slide6.shapes.add_textbox(Inches(8.4), Inches(1.75), Inches(3.95), Inches(5.0))
    tf_flow = tb_flow.text_frame
    tf_flow.word_wrap = True

    p = tf_flow.paragraphs[0]
    r_fh = p.add_run()
    r_fh.text = "CƠ CHẾ DIRECT UPLOAD PATTERN"
    r_fh.font.name = "Arial"
    r_fh.font.size = Pt(12)
    r_fh.font.bold = True
    r_fh.font.color.rgb = PRIMARY_BLUE

    steps = [
        ("1. Yêu cầu tải lên:", "Client gửi metadata lên Backend API (tên tệp, MIME, dung lượng)."),
        ("2. Cấp quyền tạm thời:", "Backend xác thực RBAC và sinh Pre-signed URL có chữ ký mã hóa (hạn 15 phút)."),
        ("3. Tải trực tiếp lên Cloud:", "Client đẩy tệp trực tiếp lên S3/Storage, bypass hoàn toàn Backend."),
        ("4. Xử lý không đồng bộ:", "S3 Event kích hoạt Lambda/Cloud Function sinh Thumbnail, OCR văn bản."),
        ("5. Phân phối toàn cầu:", "Người dùng đọc tài liệu với tốc độ cao qua CDN Edge Caching.")
    ]
    for s_title, s_desc in steps:
        p_s = tf_flow.add_paragraph()
        r_st = p_s.add_run()
        r_st.text = f"⚡ {s_title} "
        r_st.font.name = "Arial"
        r_st.font.size = Pt(9.5)
        r_st.font.bold = True
        r_st.font.color.rgb = ACCENT_TEAL
        p_s.space_before = Pt(6)

        run = p_s.add_run()
        run.text = s_desc
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 7: Security, Cost & Performance Evaluation
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "5. ĐÁNH GIÁ 3 TRỤ CỘT: BẢO MẬT, CHI PHÍ VÀ HIỆU SUẤT")

    eval_pillars = [
        ("🛡️ AN TOÀN BẢO MẬT", "Chuẩn Quốc Tế & Chống Ransomware", "• Mã hóa kép At-Rest (AES-256) & In-Transit (TLS 1.3).\n• Liên kết chia sẻ tạm thời Pre-signed URL (TTL ngắn).\n• Object Lock (WORM) bảo vệ chống xóa/ghi đè mã độc.\n• Phân quyền chi tiết tới từng tệp/thư mục (IAM & Security Rules).\n• Giám sát toàn bộ vết truy cập tệp 24/7."),
        ("💰 TỐI ƯU CHI PHÍ (TCO)", "Chuyển Đổi CapEx sang OpEx", "• Đầu tư ban đầu (CapEx) bằng 0: Không cần mua máy chủ.\n• Mô hình Pay-As-You-Go: Chỉ trả tiền theo dung lượng dùng thật.\n• Vòng đời tự động (Lifecycle Rules): Tự động chuyển tài liệu cũ sang Glacier Deep Archive, tiết kiệm 70-90% chi phí lưu trữ.\n• Tiết kiệm 40% TCO tổng thể trong 3 năm."),
        ("🚀 HIỆU SUẤT VƯỢT TRỘI", "Bền Bỉ 11 Số 9 & CDN Toàn Cầu", "• Độ bền dữ liệu 99.999999999% phân tán đa trung tâm dữ liệu.\n• Tự động co giãn (Auto-scaling) đáp ứng hàng vạn sinh viên thi cử.\n• CDN Edge Caching giảm 75% độ trễ tải tài liệu từ xa.\n• Backend giải phóng 90% tải CPU và băng thông mạng nhờ cơ chế Direct Upload.")
    ]

    for idx, (title, sub, details) in enumerate(eval_pillars):
        c_left = Inches(0.8 + idx * 3.95)
        c_top = Inches(1.6)
        c_w = Inches(3.8)
        c_h = Inches(5.3)
        add_card(slide7, c_left, c_top, c_w, c_h)

        tb = slide7.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.15), c_w - Inches(0.3), c_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r_t = p1.add_run()
        r_t.text = title
        r_t.font.name = "Arial"
        r_t.font.size = Pt(11.5)
        r_t.font.bold = True
        r_t.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        r_sub = p2.add_run()
        r_sub.text = sub
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(9.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = ACCENT_TEAL
        p2.space_after = Pt(8)

        p3 = tf.add_paragraph()
        r_det = p3.add_run()
        r_det.text = details
        r_det.font.name = "Arial"
        r_det.font.size = Pt(9.5)
        r_det.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 8: Firebase Integration Solution
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "6. GIẢI PHÁP TÍCH HỢP FIREBASE (GOOGLE AUTH & STORAGE)")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.75), Inches(5.3))
    tb_fa = slide8.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.35), Inches(5.0))
    tf_fa = tb_fa.text_frame
    tf_fa.word_wrap = True

    p = tf_fa.paragraphs[0]
    r_ah = p.add_run()
    r_ah.text = "🔑 FIREBASE AUTHENTICATION (GOOGLE SIGN-IN)"
    r_ah.font.name = "Arial"
    r_ah.font.size = Pt(12)
    r_ah.font.bold = True
    r_ah.font.color.rgb = PRIMARY_BLUE

    fa_desc = [
        ("Xác thực một chạm (OAuth 2.0):", "Sinh viên đăng nhập an toàn bằng tài khoản Google trường (@e.tlu.edu.vn / @gmail.com) mà không cần ghi nhớ thêm mật khẩu."),
        ("Thư viện Flutter chuẩn:", "firebase_auth kết hợp google_sign_in cung cấp trải nghiệm mượt mà trên cả Android, iOS và Web."),
        ("Bảo mật cao cấp:", "Tự động quản lý ID Token, Refresh Token, cơ chế đa yếu tố (MFA) được Google bảo đảm an toàn."),
        ("Tích hợp phân quyền:", "Token JWT của Firebase chứa User ID (request.auth.uid) được dùng trực tiếp để phân quyền truy cập trong Security Rules của Storage.")
    ]
    for ft, fd in fa_desc:
        p_fa = tf_fa.add_paragraph()
        r_ft = p_fa.add_run()
        r_ft.text = f"🔸 {ft} "
        r_ft.font.name = "Arial"
        r_ft.font.size = Pt(9.5)
        r_ft.font.bold = True
        r_ft.font.color.rgb = AMBER_FIREBASE
        p_fa.space_before = Pt(6)

        run = p_fa.add_run()
        run.text = fd
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = TEXT_DARK

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.3))
    tb_fs = slide8.shapes.add_textbox(Inches(7.0), Inches(1.75), Inches(5.35), Inches(5.0))
    tf_fs = tb_fs.text_frame
    tf_fs.word_wrap = True

    p_fs = tf_fs.paragraphs[0]
    r_sh = p_fs.add_run()
    r_sh.text = "📦 CLOUD STORAGE FOR FIREBASE"
    r_sh.font.name = "Arial"
    r_sh.font.size = Pt(12)
    r_sh.font.bold = True
    r_sh.font.color.rgb = PRIMARY_BLUE

    fs_desc = [
        ("Lưu trữ tệp nhị phân tin cậy:", "Xây dựng trên nền tảng Google Cloud Storage, khả năng co giãn vô hạn, lưu trữ an toàn các tệp PDF, DOCX, Slide bài giảng."),
        ("Khả năng tiếp tục tải (Resumable):", "Tự động tạm dừng và tải tiếp khi mạng chập chờn trên thiết bị di động, tối ưu hóa trải nghiệm sinh viên."),
        ("Security Rules linh hoạt:", "Kiểm soát truy cập dựa trên User ID và vai trò:\nallow read: if request.auth != null;\nallow write: if request.auth.uid == userId;"),
        ("Tải xuống trực tiếp:", "Sinh Download URL công khai hoặc liên kết có chữ ký truy cập qua CDN của Google với độ trễ thấp.")
    ]
    for st, sd in fs_desc:
        p_s = tf_fs.add_paragraph()
        r_st = p_s.add_run()
        r_st.text = f"🔹 {st} "
        r_st.font.name = "Arial"
        r_st.font.size = Pt(9.5)
        r_st.font.bold = True
        r_st.font.color.rgb = ACCENT_TEAL
        p_s.space_before = Pt(6)

        run = p_s.add_run()
        run.text = sd
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 9: Step-by-Step Firebase Flutter Setup Guide
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "7. QUY TRÌNH THIẾT LẬP FIREBASE VỚI TÀI KHOẢN NHÓM")

    steps_setup = [
        ("BƯỚC 1: KHỞI TẠO DỰ ÁN", "Tạo Project Console", "• Truy cập console.firebase.google.com bằng tài khoản nhóm.\n• Tạo dự án mới cashew-study-docs.\n• Bật Google Analytics (tùy chọn)."),
        ("BƯỚC 2: CÀI ĐẶT CLI", "Firebase CLI & FlutterFire", "• Cài đặt Firebase Tools: npm install -g firebase-tools.\n• Đăng nhập: firebase login.\n• Kích hoạt: dart pub global activate flutterfire_cli."),
        ("BƯỚC 3: CẤU HÌNH TỰ ĐỘNG", "Sinh tệp firebase_options", "• Tại thư mục app, chạy lệnh: flutterfire configure.\n• Tự động đăng ký app Android, iOS, Web và sinh mã lib/firebase_options.dart."),
        ("BƯỚC 4: THÊM DEPENDENCIES", "Cài đặt gói trong pubspec", "• Thêm vào pubspec.yaml:\nfirebase_core: ^3.6.0\nfirebase_auth: ^5.3.1\ngoogle_sign_in: ^6.2.1\nfirebase_storage: ^12.3.2"),
        ("BƯỚC 5: KHỞI TẠO CODE", "Khởi chạy trong main.dart", "• Bổ sung vào hàm main():\nWidgetsFlutterBinding.ensureInitialized();\nawait Firebase.initializeApp(\n  options: DefaultFirebaseOptions.currentPlatform);")
    ]

    for idx, (title, sub, details) in enumerate(steps_setup):
        c_left = Inches(0.8 + idx * 2.37)
        c_top = Inches(1.6)
        c_w = Inches(2.25)
        c_h = Inches(5.3)
        add_card(slide9, c_left, c_top, c_w, c_h)

        tb = slide9.shapes.add_textbox(c_left + Inches(0.1), c_top + Inches(0.12), c_w - Inches(0.2), c_h - Inches(0.25))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.bold = True
        r2.font.color.rgb = ACCENT_TEAL
        p2.space_after = Pt(6)

        p3 = tf.add_paragraph()
        r3 = p3.add_run()
        r3.text = details
        r3.font.name = "Arial"
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 10: Team Work Division for Next Assignment
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "8. BẢNG PHÂN CHIA CÔNG VIỆC BÀI TẬP TIẾP THEO (NHÓM 16)")

    team_tasks = [
        ("1. NGUYỄN VĂN HUỲNH (NHÓM TRƯỞNG)", "Quản trị Cloud, Security & Tích hợp Kho lưu trữ", "• Khởi tạo Firebase Project & cấu hình OAuth Client ID, SHA-1 fingerprint.\n• Thiết lập dịch vụ Firebase Storage và viết Security Rules kiểm soát quyền truy cập.\n• Triển khai lớp Storage Service trong Flutter (Tải lên/tải xuống tệp bài giảng, đề thi).\n• Quản lý Git Flow, review pull request và hợp nhất vào nhánh main.\n• Nhánh phụ trách: huynh-cloud-storage"),
        ("2. LÊ ANH TUẤN (THÀNH VIÊN)", "Xác thực Người dùng Google Sign-In & Auth State", "• Cài đặt và cấu hình thư viện firebase_auth & google_sign_in.\n• Triển khai luồng đăng nhập một chạm bằng tài khoản Google trường cho sinh viên.\n• Xây dựng màn hình đăng nhập (Login View), quản lý phiên làm việc AuthStateProvider.\n• Viết kịch bản kiểm thử luồng đăng nhập/đăng xuất và ghi nhận log kiểm thử.\n• Nhánh phụ trách: letuan-google-auth"),
        ("3. NGUYỄN TRUNG KIÊN (THÀNH VIÊN)", "Thiết kế Giao diện UI/UX & Đồng bộ Trạng thái Tải tệp", "• Thiết kế giao diện hiển thị người dùng (Avatar, thông tin tài khoản Google đăng nhập).\n• Thiết kế thanh tiến trình tải lên/tải xuống tệp (Upload/Download Progress Bar).\n• Tùy biến thông báo trạng thái đồng bộ Cloud (Cloud Sync Badge, Offline indicator).\n• Tối ưu giao diện xem trước tài liệu và chụp ảnh minh chứng kiểm thử.\n• Nhánh phụ trách: kien-cloud-ui"),
        ("4. TRẦN ANH TUẤN (THÀNH VIÊN)", "Xử lý Dữ liệu Ngoại tuyến, Đồng bộ Cache & Kiểm toán", "• Xây dựng cơ chế Local Cache kết hợp Cloud: Đọc dữ liệu từ SQLite khi offline.\n• Xử lý đồng bộ dữ liệu hai chiều (Cloud Firestore / Metadata Sync) khi có mạng trở lại.\n• Kiểm tra tính toàn vẹn tệp (Checksum MD5/SHA-256) và cập nhật bảng delete_logs.\n• Kiểm thử tải đồng thời nhiều tệp và đánh giá hiệu năng ứng dụng.\n• Nhánh phụ trách: trantuan-offline-sync")
    ]

    for idx, (name, role, tasks) in enumerate(team_tasks):
        row = idx // 2
        col = idx % 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.6 + row * 2.65)
        c_w = Inches(5.75)
        c_h = Inches(2.5)
        add_card(slide10, c_left, c_top, c_w, c_h)

        tb = slide10.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.1), c_w - Inches(0.3), c_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = f"👤 {name}"
        r1.font.name = "Arial"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = f"🎯 Vai trò: {role}"
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = ACCENT_TEAL
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        r3 = p3.add_run()
        r3.text = tasks
        r3.font.name = "Arial"
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 11: Summary & Conclusion
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "9. TỔNG KẾT VÀ KẾT LUẬN TOÀN DIỆN")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.3))
    tb_end = slide11.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.9))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    r_eh = p.add_run()
    r_eh.text = "HOÀN THÀNH TOÀN DIỆN 7/7 MỤC CHECKLIST ĐỀ BÀI YÊU CẦU"
    r_eh.font.name = "Arial"
    r_eh.font.size = Pt(14)
    r_eh.font.bold = True
    r_eh.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(14)

    concl_points = [
        ("Khẳng định giá trị chuyển đổi Cloud:", "Chuyển đổi từ On-Premises sang Cloud không chỉ mở rộng dung lượng không giới hạn, mà còn gia tăng vượt trội tính an toàn, bảo mật, khả năng truy cập mọi lúc mọi nơi cho sinh viên và giảng viên."),
        ("Hệ sinh thái đối sánh chuyên sâu:", "Kết hợp mô hình lý thuyết tổng thể chuẩn AWS (S3, RDS, CloudFront, Lambda) với giải pháp thực nghiệm nhanh gọn, mạnh mẽ của Firebase (Google Sign-In & Cloud Storage) tạo nên bộ giải pháp công nghệ toàn diện."),
        ("Sẵn sàng triển khai thực nghiệm:", "Các bước thiết lập Firebase cho Flutter đã được chuẩn hóa rõ ràng. Nhóm 16 đã phân công cụ thể từng đầu việc, từng nhánh Git cho 4 thành viên, sẵn sàng bắt tay vào lập trình bài tập tiếp theo."),
        ("Hồ sơ nộp bài đầy đủ:", "Bao gồm Slide thuyết trình PPTX chuyên nghiệp, Báo cáo kỹ thuật chi tiết (.DOCX và .MD), cùng tài liệu README.MD được cập nhật hoàn chỉnh trên GitHub.")
    ]

    for ct, cd in concl_points:
        p_c = tf_end.add_paragraph()
        r_ct = p_c.add_run()
        r_ct.text = f"✨ {ct} "
        r_ct.font.name = "Arial"
        r_ct.font.size = Pt(10.5)
        r_ct.font.bold = True
        r_ct.font.color.rgb = PRIMARY_BLUE
        p_c.space_before = Pt(8)

        run = p_c.add_run()
        run.text = cd
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = TEXT_DARK

    # Save presentation
    prs_filename = "Slide_Tich_Hop_Cloud_Firebase_DMS.pptx"
    prs.save(prs_filename)
    print(f"Presentation PPTX successfully created: {prs_filename}")

if __name__ == '__main__':
    create_presentation()
